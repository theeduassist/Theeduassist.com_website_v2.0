import openpyxl

for fname in ['TheEduAssist_Content_Management_System.xlsx', 'TheEduAssist_SEO_Operations_Tracker.xlsx']:
    wb = openpyxl.load_workbook(fname, read_only=True)
    print(f"=== {fname} ===")
    print(f"Sheets: {wb.sheetnames}")
    for name in wb.sheetnames[:3]:
        sheet = wb[name]
        headers = [cell for cell in next(sheet.iter_rows(values_only=True), [])]
        print(f"  Sheet '{name}' (first 10 headers): {headers[:10]}")
